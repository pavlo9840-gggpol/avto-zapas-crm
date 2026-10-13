# -*- coding: utf-8 -*-
"""
app.py
------
Веб-версія обліку шин, акумуляторів та мастил (Flask + SQLite).
Облікові записи адміністраторів і продавців зберігаються у SQLite.
"""

import os
import re
import secrets
import json
import csv
import io
import math
from functools import wraps
from datetime import datetime, timedelta

from flask import (
    Flask, render_template, request, redirect, url_for, session, flash, g, Response
)
from werkzeug.security import generate_password_hash, check_password_hash
from markupsafe import Markup, escape

try:
    import openpyxl
except ImportError:
    openpyxl = None

import db as dbmod
from constants import (
    ALL_CATEGORIES, CATEGORY_LABELS, CATEGORY_TYPE, CATEGORY_GROUPS,
    CATEGORY_ICONS, GROUP_ICONS,
    SEASONS, CAR_WIDTHS, CAR_PROFILES, CAR_DIAMETERS, MOTO_SIZES, TIRE_BRANDS,
    LOAD_INDEXES, SPEED_INDEXES, LOAD_INDEX_TABLE, SPEED_INDEX_TABLE,
    BATTERY_BRANDS, BATTERY_TRANSPORT_TYPES,
    BATTERY_POLARITY_OPTIONS, BATTERY_CASE_TYPES, OIL_BRANDS, OIL_VISCOSITY_TYPES,
    OIL_BASE_TYPES, OIL_VOLUMES, is_valid_dot_code, decode_dot_code, decode_li_si,
    format_number, COLOR_TAGS, COLOR_TAG_HEX, ALL_KNOWN_BRANDS, compare_tire_sizes,
    DOC_PLACEHOLDERS_HELP, DEFAULT_RECEIPT_TEMPLATE, DEFAULT_INVOICE_TEMPLATE,
    DEFAULT_GOODS_INVOICE_TEMPLATE,
)

ALLOWED_ICON_EXT = {"png", "ico"}
ROLE_LABELS = {"admin": "Адміністратор", "seller": "Продавець"}

IMPORT_COLUMNS = {
    "tire": [
        ("Бренд", "brand"), ("Модель", "tread_model"), ("Розмір", "size"), ("Сезон", "season"),
        ("LI", "load_index"), ("SI", "speed_index"), ("DOT", "dot_code"), ("Мітка", "color_tag"),
        ("Собівартість", "cost_price"), ("Ціна", "price"), ("Кількість", "quantity"),
    ],
    "battery": [
        ("Бренд", "brand"), ("Тип", "transport_type"), ("Ємність (Ah)", "capacity_ah"),
        ("Пусковий струм (A)", "cold_crank_amps"), ("Полярність", "polarity"), ("Корпус", "case_type"),
        ("Мітка", "color_tag"), ("Собівартість", "cost_price"), ("Ціна", "price"), ("Кількість", "quantity"),
    ],
    "oil": [
        ("Бренд", "brand"), ("Назва", "title"), ("В'язкість", "viscosity"), ("Основа", "base_type"),
        ("Обʼєм", "volume"), ("Мітка", "color_tag"), ("Собівартість", "cost_price"),
        ("Ціна", "price"), ("Кількість", "quantity"),
    ],
}

app = Flask(__name__)
app.config["DB_PATH"] = os.environ.get("STOCK_DB_PATH", dbmod.DEFAULT_DB_PATH)
app.config["MAX_CONTENT_LENGTH"] = 2 * 1024 * 1024  # 2 MB - достатньо для маленької іконки бренду
app.config["DONATION_URL"] = os.environ.get("DONATION_URL", "").strip()
dbmod.register(app)

BRAND_ICON_DIR = os.path.join(app.static_folder, "brand_icons")
os.makedirs(BRAND_ICON_DIR, exist_ok=True)


@app.errorhandler(413)
def too_large(_e):
    flash("Файл завеликий (максимум 2 МБ). Виберіть менший PNG/ICO.", "error")
    return redirect(request.referrer or url_for("brand_icons_view"))


def _slugify(text):
    text = re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower()
    return text or "brand"


def _get_or_create_secret_key():
    """Секретний ключ сесій зберігаємо в БД, щоб сесії не злітали при перезапуску сервера."""
    with app.app_context():
        db = dbmod.get_db()
        row = db.execute("SELECT value FROM settings WHERE key = 'secret_key'").fetchone()
        if row:
            return row["value"]
        key = secrets.token_hex(32)
        db.execute("INSERT INTO settings (key, value) VALUES ('secret_key', ?)", (key,))
        db.commit()
        return key


app.secret_key = os.environ.get("FLASK_SECRET_KEY") or _get_or_create_secret_key()

TIRE_KEY_FIELDS = ("brand", "size", "season", "dot_code", "tread_model", "load_index", "speed_index")
BATTERY_KEY_FIELDS = ("brand", "transport_type", "capacity_ah", "cold_crank_amps", "polarity", "case_type")
OIL_KEY_FIELDS = ("brand", "title", "viscosity", "base_type", "volume")

KEY_FIELDS_BY_TYPE = {"tire": TIRE_KEY_FIELDS, "battery": BATTERY_KEY_FIELDS, "oil": OIL_KEY_FIELDS}

ALL_FIELDS = (
    "brand", "tread_model", "size", "season", "load_index", "speed_index", "dot_code",
    "transport_type", "capacity_ah", "cold_crank_amps", "polarity", "case_type",
    "title", "viscosity", "base_type", "volume",
)


# ------------------------------------------------------------------
#                          Аутентифікація
# ------------------------------------------------------------------

def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("authed"):
            return redirect(url_for("login", next=request.path))
        return view(*args, **kwargs)
    return wrapped


def role_required(*roles):
    def decorator(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            if not session.get("authed"):
                return redirect(url_for("login", next=request.path))
            if session.get("role") not in roles:
                flash("Ця дія доступна лише адміністратору.", "error")
                return redirect(url_for("dashboard"))
            return view(*args, **kwargs)
        return wrapped
    return decorator


@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("authed"):
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        username = (request.form.get("username", "") or "").strip()
        password = request.form.get("password", "")
        db = dbmod.get_db()
        user = db.execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()
        if user and check_password_hash(user["password_hash"], password):
            session.clear()
            session["authed"] = True
            session["user_id"] = user["id"]
            session["username"] = user["username"]
            session["role"] = user["role"]
            session.permanent = True
            nxt = request.args.get("next") or url_for("dashboard")
            return redirect(nxt)
        flash("Невірний логін або пароль.", "error")
    return render_template("login.html", donation_url=app.config["DONATION_URL"])


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


# ------------------------------------------------------------------
#                          Допоміжне
# ------------------------------------------------------------------

def _category_or_404(category):
    if category not in ALL_CATEGORIES:
        from flask import abort
        abort(404)
    return category


def _get_threshold(category):
    db = dbmod.get_db()
    row = db.execute("SELECT min_qty FROM thresholds WHERE category = ?", (category,)).fetchone()
    return row["min_qty"] if row else 3


def _set_threshold(category, value):
    db = dbmod.get_db()
    db.execute(
        "INSERT INTO thresholds (category, min_qty) VALUES (?, ?) "
        "ON CONFLICT(category) DO UPDATE SET min_qty = excluded.min_qty",
        (category, value),
    )
    db.commit()


def _find_matching_item(db, category, values):
    """Шукає позицію на складі з такими самими ключовими атрибутами (щоб не плодити дублі)."""
    cat_type = CATEGORY_TYPE[category]
    key_fields = KEY_FIELDS_BY_TYPE[cat_type]
    where = " AND ".join(f"{f} = ?" for f in key_fields)
    params = [values.get(f, "") or "" for f in key_fields]
    row = db.execute(
        f"SELECT * FROM stock_items WHERE category = ? AND {where}",
        [category] + params,
    ).fetchone()
    return row


def _item_label(category, item):
    cat_type = CATEGORY_TYPE[category]
    if cat_type == "tire":
        tread = f" {item['tread_model']}" if item["tread_model"] else ""
        dot = f" · DOT {item['dot_code']}" if item["dot_code"] else ""
        return f"{item['brand']}{tread} · {item['size']} · {item['season']}{dot}"
    if cat_type == "battery":
        return f"{item['brand']} · {item['transport_type']} · {format_number(item['capacity_ah'])} Ah / {format_number(item['cold_crank_amps'])} A"
    title = f" {item['title']}" if item["title"] else ""
    return f"{item['brand']}{title} · {item['viscosity']} ({item['volume']})"


def _add_history(db, category, operation, item_values, quantity, customer="",
                  customer_id=None, sale_id="", username=""):
    fields = ", ".join(ALL_FIELDS)
    placeholders = ", ".join("?" for _ in ALL_FIELDS)
    values = [item_values.get(f, "") or "" for f in ALL_FIELDS]
    db.execute(
        f"INSERT INTO history (category, operation, datetime, {fields}, cost_price, price, quantity, "
        f"customer, customer_id, sale_id, username) "
        f"VALUES (?, ?, ?, {placeholders}, ?, ?, ?, ?, ?, ?, ?)",
        [category, operation, dbmod.now_str()] + values +
        [item_values.get("cost_price", 0) or 0, item_values.get("price", 0) or 0, quantity,
         customer, customer_id, sale_id, username or session.get("username", "")],
    )


def _get_brand_icons(db):
    rows = db.execute(
        "SELECT brand, icon, icon_file FROM brand_icons WHERE icon != '' OR icon_file != ''"
    ).fetchall()
    return {r["brand"]: {"icon": r["icon"], "icon_file": r["icon_file"]} for r in rows}


def brand_icon_html(brand, icons_map):
    info = icons_map.get(brand) if icons_map else None
    if not info:
        return ""
    if info.get("icon_file"):
        url = url_for("static", filename="brand_icons/" + info["icon_file"])
        return Markup(f'<img src="{url}" class="brand-icon-img" alt="">')
    if info.get("icon"):
        return Markup(f'<span class="brand-icon-emoji">{escape(info["icon"])}</span>')
    return ""


app.jinja_env.globals.update(
    CATEGORY_LABELS=CATEGORY_LABELS, CATEGORY_TYPE=CATEGORY_TYPE, CATEGORY_GROUPS=CATEGORY_GROUPS,
    format_number=format_number, decode_dot_code=decode_dot_code, decode_li_si=decode_li_si,
    COLOR_TAGS=COLOR_TAGS, COLOR_TAG_HEX=COLOR_TAG_HEX, brand_icon_html=brand_icon_html,
    ROLE_LABELS=ROLE_LABELS, CATEGORY_ICONS=CATEGORY_ICONS, GROUP_ICONS=GROUP_ICONS,
)


@app.context_processor
def inject_nav_low_stock():
    """Одним легким запитом рахує, скільки позицій 'закінчується' у кожній категорії -
    використовується для бейджів у бічній навігації."""
    if not session.get("authed"):
        return {"nav_low_stock": {}}
    db = dbmod.get_db()
    rows = db.execute(
        "SELECT category, COUNT(*) AS cnt FROM stock_items si "
        "WHERE quantity - COALESCE((SELECT SUM(quantity) FROM reservations r "
        "WHERE r.item_id = si.id AND r.status = 'active'), 0) "
        "<= COALESCE((SELECT min_qty FROM thresholds t WHERE t.category = si.category), 3) "
        "GROUP BY category"
    ).fetchall()
    return {"nav_low_stock": {r["category"]: r["cnt"] for r in rows}}


@app.route("/search")
@login_required
def search_view():
    q = (request.args.get("q", "") or "").strip()
    results = []
    if q:
        db = dbmod.get_db()
        like = f"%{q}%"
        for category in ALL_CATEGORIES:
            cat_type = CATEGORY_TYPE[category]
            if cat_type == "tire":
                cols = "brand LIKE ? OR tread_model LIKE ? OR size LIKE ? OR dot_code LIKE ?"
                params = [like, like, like, like]
            elif cat_type == "battery":
                cols = "brand LIKE ? OR transport_type LIKE ?"
                params = [like, like]
            else:
                cols = "brand LIKE ? OR title LIKE ? OR viscosity LIKE ?"
                params = [like, like, like]
            rows = db.execute(
                f"SELECT * FROM stock_items WHERE category = ? AND ({cols}) ORDER BY brand LIMIT 50",
                [category] + params,
            ).fetchall()
            for r in rows:
                results.append((category, r))
    return render_template("search.html", q=q, results=results, item_label=_item_label)


# ------------------------------------------------------------------
#                          Дашборд
# ------------------------------------------------------------------

@app.route("/")
@login_required
def dashboard():
    db = dbmod.get_db()
    is_admin = session.get("role") == "admin"

    total_value = db.execute("SELECT COALESCE(SUM(price * quantity), 0) AS v FROM stock_items").fetchone()["v"]

    low_items = []
    for category in ALL_CATEGORIES:
        threshold = _get_threshold(category)
        if threshold <= 0:
            continue
        rows = db.execute(
            "SELECT si.*, si.quantity - COALESCE((SELECT SUM(quantity) FROM reservations r "
            "WHERE r.item_id = si.id AND r.status = 'active'), 0) AS available_quantity "
            "FROM stock_items si WHERE category = ? AND quantity - COALESCE((SELECT SUM(quantity) "
            "FROM reservations r WHERE r.item_id = si.id AND r.status = 'active'), 0) <= ? "
            "ORDER BY available_quantity ASC",
            (category, threshold),
        ).fetchall()
        for r in rows:
            low_items.append((category, r))
    low_items.sort(key=lambda t: t[1]["available_quantity"])

    today = dbmod.now_str()[:10]
    sales_row = db.execute(
        "SELECT COUNT(*) AS cnt, COALESCE(SUM(price * quantity), 0) AS revenue "
        "FROM history h WHERE operation = 'Продаж' AND datetime LIKE ? "
        "AND NOT EXISTS (SELECT 1 FROM sales s WHERE s.sale_id = h.sale_id AND s.canceled_at IS NOT NULL)",
        (f"{today}%",),
    ).fetchone()

    total_potential_profit = today_profit = None
    if is_admin:
        total_potential_profit = db.execute(
            "SELECT COALESCE(SUM((price - cost_price) * quantity), 0) AS v FROM stock_items"
        ).fetchone()["v"]
        today_profit = db.execute(
            "SELECT COALESCE(SUM((price - cost_price) * quantity), 0) AS v "
            "FROM history h WHERE operation = 'Продаж' AND datetime LIKE ? "
            "AND NOT EXISTS (SELECT 1 FROM sales s WHERE s.sale_id = h.sale_id AND s.canceled_at IS NOT NULL)",
            (f"{today}%",),
        ).fetchone()["v"]

    recent = db.execute("SELECT * FROM history ORDER BY id DESC LIMIT 15").fetchall()
    brand_icons = _get_brand_icons(db)

    return render_template(
        "dashboard.html",
        total_value=total_value, low_items=low_items[:20], low_count=len(low_items),
        sales_count=sales_row["cnt"], sales_revenue=sales_row["revenue"], recent=recent,
        item_label=_item_label, brand_icons=brand_icons,
        total_potential_profit=total_potential_profit, today_profit=today_profit,
    )


@app.route("/reports")
@login_required
def reports_view():
    db = dbmod.get_db()
    period = request.args.get("period", "week")
    if period not in ("week", "month"):
        period = "week"
    today = datetime.now().date()
    if period == "week":
        start = today - timedelta(days=today.weekday())
        end = start + timedelta(days=7)
        period_label = "Цей тиждень"
    else:
        start = today.replace(day=1)
        end = today.replace(year=today.year + 1, month=1, day=1) if today.month == 12 else today.replace(month=today.month + 1, day=1)
        period_label = "Цей місяць"
    start_at = start.strftime("%Y-%m-%d 00:00:00")
    end_at = end.strftime("%Y-%m-%d 00:00:00")
    sale_filter = (
        "h.operation = 'Продаж' AND h.datetime >= ? AND h.datetime < ? "
        "AND NOT EXISTS (SELECT 1 FROM sales s WHERE s.sale_id = h.sale_id AND s.canceled_at IS NOT NULL)"
    )
    summary = db.execute(
        "SELECT COUNT(DISTINCT h.sale_id) AS sales, COALESCE(SUM(h.price * h.quantity), 0) AS revenue, "
        "COALESCE(SUM((h.price - h.cost_price) * h.quantity), 0) AS profit "
        f"FROM history h WHERE {sale_filter}", (start_at, end_at),
    ).fetchone()
    by_category = db.execute(
        "SELECT h.category, COUNT(DISTINCT h.sale_id) AS sales, SUM(h.quantity) AS quantity, "
        "SUM(h.price * h.quantity) AS revenue, SUM((h.price - h.cost_price) * h.quantity) AS profit "
        f"FROM history h WHERE {sale_filter} GROUP BY h.category ORDER BY revenue DESC", (start_at, end_at),
    ).fetchall()
    by_seller = db.execute(
        "SELECT COALESCE(NULLIF(h.username, ''), '—') AS username, COUNT(DISTINCT h.sale_id) AS sales, "
        "SUM(h.price * h.quantity) AS revenue, SUM((h.price - h.cost_price) * h.quantity) AS profit "
        f"FROM history h WHERE {sale_filter} GROUP BY h.username ORDER BY revenue DESC", (start_at, end_at),
    ).fetchall()
    return render_template(
        "reports.html", period=period, period_label=period_label, summary=summary,
        by_category=by_category, by_seller=by_seller, is_admin=session.get("role") == "admin",
    )


# ------------------------------------------------------------------
#                          Список складу
# ------------------------------------------------------------------

@app.route("/stock/<category>")
@login_required
def stock_list(category):
    _category_or_404(category)
    cat_type = CATEGORY_TYPE[category]
    db = dbmod.get_db()
    items = db.execute(
        "SELECT * FROM stock_items WHERE category = ? ORDER BY brand, size, title", (category,)
    ).fetchall()
    items_with_availability = []
    for item in items:
        reserved = db.execute(
            "SELECT COALESCE(SUM(quantity), 0) AS quantity FROM reservations "
            "WHERE item_id = ? AND status = 'active'", (item["id"],),
        ).fetchone()["quantity"]
        items_with_availability.append({**dict(item), "available_quantity": item["quantity"] - reserved})
    threshold = _get_threshold(category)

    total_positions = len(items)
    total_quantity = sum(item["quantity"] for item in items)

    filter_options = {}
    if cat_type == "tire":
        filter_options["season"] = sorted({i["season"] for i in items if i["season"]})
        filter_options["size"] = sorted({i["size"] for i in items if i["size"]})
    elif cat_type == "battery":
        filter_options["transport_type"] = sorted({i["transport_type"] for i in items if i["transport_type"]})
    else:
        filter_options["viscosity"] = sorted({i["viscosity"] for i in items if i["viscosity"]})
        filter_options["base_type"] = sorted({i["base_type"] for i in items if i["base_type"]})

    customer_names = [r["name"] for r in db.execute("SELECT name FROM customers ORDER BY name")]

    return render_template(
        "stock.html", category=category, items=items_with_availability, threshold=threshold,
        seasons=SEASONS, total_positions=total_positions, total_quantity=total_quantity,
        filter_options=filter_options, brand_icons=_get_brand_icons(db), customer_names=customer_names,
    )


@app.route("/stock/<category>/threshold", methods=["POST"])
@role_required("admin")
def set_threshold(category):
    _category_or_404(category)
    try:
        value = max(0, int(request.form.get("min_qty", 0)))
    except ValueError:
        value = 0
    _set_threshold(category, value)
    return redirect(url_for("stock_list", category=category))


# ------------------------------------------------------------------
#                          Додати / редагувати
# ------------------------------------------------------------------

def _form_options(category):
    cat_type = CATEGORY_TYPE[category]
    opts = {"cat_type": cat_type, "seasons": SEASONS}
    if cat_type == "tire":
        opts.update(
            car_widths=CAR_WIDTHS, car_profiles=CAR_PROFILES, car_diameters=CAR_DIAMETERS,
            moto_sizes=MOTO_SIZES, brands=TIRE_BRANDS,
            load_indexes=LOAD_INDEXES, speed_indexes=SPEED_INDEXES,
            is_car=(category == "car"),
            load_index_table_json=json.dumps(LOAD_INDEX_TABLE),
            speed_index_table_json=json.dumps(SPEED_INDEX_TABLE),
        )
    elif cat_type == "battery":
        opts.update(
            brands=BATTERY_BRANDS, transport_types=BATTERY_TRANSPORT_TYPES,
            polarity_options=BATTERY_POLARITY_OPTIONS, case_types=BATTERY_CASE_TYPES,
        )
    else:
        opts.update(
            brands=OIL_BRANDS, viscosity_types=OIL_VISCOSITY_TYPES,
            base_types=OIL_BASE_TYPES, volumes=OIL_VOLUMES,
        )
    return opts


def _read_item_form(category):
    cat_type = CATEGORY_TYPE[category]
    values = {f: (request.form.get(f, "") or "").strip() for f in ALL_FIELDS}
    color_tag = (request.form.get("color_tag", "") or "").strip()
    valid_color_codes = {code for code, _label, _hex in COLOR_TAGS}
    if color_tag not in valid_color_codes:
        color_tag = ""

    if cat_type == "tire" and category == "car":
        w, p, d = request.form.get("width", ""), request.form.get("profile", ""), request.form.get("diameter", "")
        values["size"] = f"{w}/{p} {d}"

    if cat_type == "battery":
        try:
            values["capacity_ah"] = float((request.form.get("capacity_ah") or "0").replace(",", "."))
            values["cold_crank_amps"] = float((request.form.get("cold_crank_amps") or "0").replace(",", "."))
        except ValueError:
            values["capacity_ah"] = 0
            values["cold_crank_amps"] = 0

    try:
        price = float((request.form.get("price") or "0").replace(",", "."))
    except ValueError:
        price = 0.0
    try:
        cost_price = float((request.form.get("cost_price") or "0").replace(",", "."))
    except ValueError:
        cost_price = 0.0
    try:
        quantity = int(request.form.get("quantity") or "0")
    except ValueError:
        quantity = 0

    values["price"] = price
    values["cost_price"] = cost_price
    values["color_tag"] = color_tag
    return values, quantity


@app.route("/stock/<category>/add", methods=["GET", "POST"])
@role_required("admin")
def item_add(category):
    _category_or_404(category)
    cat_type = CATEGORY_TYPE[category]

    if request.method == "POST":
        values, quantity = _read_item_form(category)
        if quantity <= 0:
            flash("Кількість має бути більшою за нуль.", "error")
            return render_template("item_form.html", category=category, item=values, is_edit=False, **_form_options(category))
        if cat_type == "tire" and not is_valid_dot_code(values.get("dot_code", "")):
            flash("Некоректний DOT код (3 або 4 цифри).", "error")
            return render_template("item_form.html", category=category, item=values, is_edit=False, **_form_options(category))

        db = dbmod.get_db()
        existing = _find_matching_item(db, category, values)
        if existing:
            db.execute(
                "UPDATE stock_items SET quantity = quantity + ?, price = ?, cost_price = ? WHERE id = ?",
                (quantity, values["price"], values["cost_price"], existing["id"]),
            )
        else:
            fields = ", ".join(ALL_FIELDS)
            placeholders = ", ".join("?" for _ in ALL_FIELDS)
            vals = [values.get(f, "") for f in ALL_FIELDS]
            db.execute(
                f"INSERT INTO stock_items (category, {fields}, color_tag, cost_price, price, quantity) "
                f"VALUES (?, {placeholders}, ?, ?, ?, ?)",
                [category] + vals + [values["color_tag"], values["cost_price"], values["price"], quantity],
            )
        _add_history(db, category, "Прихід", values, quantity)
        db.commit()
        flash("Позицію додано на склад.", "success")
        return redirect(url_for("stock_list", category=category))

    return render_template("item_form.html", category=category, item=None, is_edit=False, **_form_options(category))


@app.route("/stock/<category>/edit/<int:item_id>", methods=["GET", "POST"])
@role_required("admin")
def item_edit(category, item_id):
    _category_or_404(category)
    cat_type = CATEGORY_TYPE[category]
    db = dbmod.get_db()
    item = db.execute("SELECT * FROM stock_items WHERE id = ? AND category = ?", (item_id, category)).fetchone()
    if item is None:
        flash("Позицію не знайдено.", "error")
        return redirect(url_for("stock_list", category=category))

    if request.method == "POST":
        values, quantity = _read_item_form(category)
        reserved = db.execute(
            "SELECT COALESCE(SUM(quantity), 0) AS quantity FROM reservations "
            "WHERE item_id = ? AND status = 'active'", (item_id,),
        ).fetchone()["quantity"]
        if quantity <= 0:
            flash("Кількість має бути більшою за нуль.", "error")
            return render_template("item_form.html", category=category, item=item, is_edit=True, **_form_options(category))
        if quantity < reserved:
            flash(f"Кількість не може бути меншою за зарезервовані {reserved} шт.", "error")
            return render_template("item_form.html", category=category, item=item, is_edit=True, **_form_options(category))
        if cat_type == "tire" and not is_valid_dot_code(values.get("dot_code", "")):
            flash("Некоректний DOT код (3 або 4 цифри).", "error")
            return render_template("item_form.html", category=category, item=item, is_edit=True, **_form_options(category))

        fields_set = ", ".join(f"{f} = ?" for f in ALL_FIELDS)
        vals = [values.get(f, "") for f in ALL_FIELDS]
        db.execute(
            f"UPDATE stock_items SET {fields_set}, color_tag = ?, cost_price = ?, price = ?, quantity = ? WHERE id = ?",
            vals + [values["color_tag"], values["cost_price"], values["price"], quantity, item_id],
        )
        _add_history(db, category, "Редагування", values, quantity)
        db.commit()
        flash("Зміни збережено.", "success")
        return redirect(url_for("stock_list", category=category))

    return render_template("item_form.html", category=category, item=item, is_edit=True, **_form_options(category))


@app.route("/stock/<category>/delete/<int:item_id>", methods=["POST"])
@role_required("admin")
def item_delete(category, item_id):
    _category_or_404(category)
    db = dbmod.get_db()
    item = db.execute("SELECT * FROM stock_items WHERE id = ? AND category = ?", (item_id, category)).fetchone()
    if item is None:
        flash("Позицію не знайдено.", "error")
        return redirect(url_for("stock_list", category=category))
    reserved = db.execute(
        "SELECT COALESCE(SUM(quantity), 0) AS quantity FROM reservations "
        "WHERE item_id = ? AND status = 'active'", (item_id,),
    ).fetchone()["quantity"]
    if reserved:
        flash(f"Спочатку скасуйте резерви цієї позиції ({reserved} шт.).", "error")
        return redirect(url_for("stock_list", category=category))
    db.execute("DELETE FROM stock_items WHERE id = ?", (item_id,))
    _add_history(db, category, "Видалено", dict(item), item["quantity"])
    db.commit()
    flash("Позицію видалено.", "success")
    return redirect(url_for("stock_list", category=category))


# ------------------------------------------------------------------
#                     Продаж (кошик з кількома позиціями)
# ------------------------------------------------------------------

def _find_or_create_customer(db, name, phone=""):
    name = (name or "").strip()
    if not name:
        return None
    row = db.execute("SELECT id FROM customers WHERE LOWER(name) = LOWER(?)", (name,)).fetchone()
    if row:
        if phone:
            db.execute("UPDATE customers SET phone = ? WHERE id = ? AND phone = ''", (phone, row["id"]))
        return row["id"]
    cur = db.execute(
        "INSERT INTO customers (name, phone, notes, created_at) VALUES (?, ?, '', ?)",
        (name, phone, dbmod.now_str()),
    )
    return cur.lastrowid


def _sale_total(items):
    return round(sum(item["price"] * quantity for _category, item, quantity in items), 2)


def _save_sale(db, items, customer_name, received, reservation_id=None):
    customer_id = _find_or_create_customer(db, customer_name)
    sale_id = secrets.token_hex(6)
    username = session.get("username", "")
    total = _sale_total(items)
    received = total if received is None else received
    db.execute(
        "INSERT INTO sales (sale_id, received, change_amount, reservation_id) VALUES (?, ?, ?, ?)",
        (sale_id, received, round(received - total, 2), reservation_id),
    )
    for category, item, qty in items:
        updated = db.execute(
            "UPDATE stock_items SET quantity = quantity - ? WHERE id = ? "
            "AND quantity - ? >= COALESCE((SELECT SUM(quantity) FROM reservations "
            "WHERE item_id = stock_items.id AND status = 'active' "
            "AND (? IS NULL OR id != ?)), 0)",
            (qty, item["id"], qty, reservation_id, reservation_id),
        )
        if not updated.rowcount:
            raise ValueError("Товар уже продано або зарезервовано в іншій операції.")
        db.execute("DELETE FROM stock_items WHERE id = ? AND quantity <= 0", (item["id"],))
        _add_history(
            db, category, "Продаж", dict(item), qty,
            customer=customer_name, customer_id=customer_id, sale_id=sale_id, username=username,
        )
    return sale_id, total, received


def _read_received_amount(raw):
    try:
        amount = float((raw or "").strip().replace(",", "."))
    except ValueError:
        return None
    if not math.isfinite(amount) or amount < 0:
        return None
    return round(amount, 2)


@app.route("/stock/<category>/sell", methods=["POST"])
@login_required
def item_sell(category):
    _category_or_404(category)
    return _process_sale(category, require_payment=False)


def _process_sale(category=None, require_payment=False):
    db = dbmod.get_db()
    customer_name = (request.form.get("customer", "") or "").strip()
    selected_ids = [k[len("sell_"):] for k in request.form if k.startswith("sell_")]
    if not selected_ids:
        flash("Виберіть хоча б одну позицію для продажу.", "error")
        return redirect(url_for("cash_view" if category is None else "stock_list", **({} if category is None else {"category": category})))

    to_sell = []
    for raw_id in selected_ids:
        try:
            item_id = int(raw_id)
        except ValueError:
            flash("Некоректно вибрано товар.", "error")
            return redirect(url_for("cash_view" if category is None else "stock_list", **({} if category is None else {"category": category})))
        if category is None:
            item = db.execute("SELECT * FROM stock_items WHERE id = ?", (item_id,)).fetchone()
        else:
            item = db.execute("SELECT * FROM stock_items WHERE id = ? AND category = ?", (item_id, category)).fetchone()
        if item is None:
            flash("Один із вибраних товарів більше не є на складі.", "error")
            return redirect(url_for("cash_view" if category is None else "stock_list", **({} if category is None else {"category": category})))
        raw_qty = request.form.get(f"qty_{raw_id}", "").strip()
        try:
            qty = int(raw_qty)
        except ValueError:
            qty = 0
        if qty <= 0:
            flash(f"Некоректна кількість для: {_item_label(item['category'], item)}", "error")
            return redirect(url_for("cash_view" if category is None else "stock_list", **({} if category is None else {"category": category})))
        reserved = db.execute(
            "SELECT COALESCE(SUM(quantity), 0) AS quantity FROM reservations "
            "WHERE item_id = ? AND status = 'active'", (item_id,),
        ).fetchone()["quantity"]
        available = item["quantity"] - reserved
        if qty > available:
            flash(f"Доступно лише {available} шт.: {_item_label(item['category'], item)}", "error")
            return redirect(url_for("cash_view" if category is None else "stock_list", **({} if category is None else {"category": category})))
        to_sell.append((item["category"], item, qty))

    total = _sale_total(to_sell)
    received = None
    if require_payment:
        received = _read_received_amount(request.form.get("received", ""))
        if received is None:
            flash("Введіть коректну суму отриманих коштів.", "error")
            return redirect(url_for("cash_view"))
        if received < total:
            flash(f"Отриманих коштів недостатньо. До сплати {total:.2f} ₴.", "error")
            return redirect(url_for("cash_view"))

    try:
        sale_id, total, received = _save_sale(db, to_sell, customer_name, received)
    except ValueError as error:
        db.rollback()
        flash(str(error), "error")
        return redirect(url_for("cash_view" if category is None else "stock_list", **({} if category is None else {"category": category})))
    db.commit()
    flash(f"Продаж оформлено. Решта: {received - total:.2f} ₴.", "success")
    return redirect(url_for("receipt_view", sale_id=sale_id))


@app.route("/cash", methods=["GET", "POST"])
@login_required
def cash_view():
    db = dbmod.get_db()
    if request.method == "POST":
        return _process_sale(require_payment=True)
    rows = db.execute("SELECT * FROM stock_items ORDER BY category, brand, size, title").fetchall()
    items = []
    for item in rows:
        reserved = db.execute(
            "SELECT COALESCE(SUM(quantity), 0) AS quantity FROM reservations "
            "WHERE item_id = ? AND status = 'active'", (item["id"],),
        ).fetchone()["quantity"]
        items.append({"item": item, "available": item["quantity"] - reserved})
    customer_names = [r["name"] for r in db.execute("SELECT name FROM customers ORDER BY name")]
    return render_template("cash.html", items=items, customer_names=customer_names, item_label=_item_label)


@app.route("/reservations", methods=["GET", "POST"])
@login_required
def reservations_view():
    db = dbmod.get_db()
    if request.method == "POST":
        try:
            item_id = int(request.form.get("item_id", ""))
            quantity = int(request.form.get("quantity", ""))
        except ValueError:
            flash("Виберіть товар і вкажіть коректну кількість.", "error")
            return redirect(url_for("reservations_view"))
        item = db.execute("SELECT * FROM stock_items WHERE id = ?", (item_id,)).fetchone()
        if item is None or quantity <= 0:
            flash("Товар не знайдено або кількість некоректна.", "error")
            return redirect(url_for("reservations_view"))
        created = db.execute(
            "INSERT INTO reservations (item_id, category, quantity, customer, phone, notes, created_at, created_by) "
            "SELECT id, category, ?, ?, ?, ?, ?, ? FROM stock_items "
            "WHERE id = ? AND quantity - COALESCE((SELECT SUM(quantity) FROM reservations "
            "WHERE item_id = stock_items.id AND status = 'active'), 0) >= ?",
            (quantity, (request.form.get("customer", "") or "").strip(),
             (request.form.get("phone", "") or "").strip(),
             (request.form.get("notes", "") or "").strip(), dbmod.now_str(),
             session.get("username", ""), item_id, quantity),
        )
        if not created.rowcount:
            flash("Товару вже недостатньо для такого резерву.", "error")
            return redirect(url_for("reservations_view"))
        db.commit()
        flash("Товар зарезервовано.", "success")
        return redirect(url_for("reservations_view"))

    stock_rows = db.execute("SELECT * FROM stock_items ORDER BY category, brand, size, title").fetchall()
    items = []
    for item in stock_rows:
        reserved = db.execute(
            "SELECT COALESCE(SUM(quantity), 0) AS quantity FROM reservations "
            "WHERE item_id = ? AND status = 'active'", (item["id"],),
        ).fetchone()["quantity"]
        items.append({**dict(item), "available_quantity": item["quantity"] - reserved})
    active = db.execute(
        "SELECT r.id AS reservation_id, r.quantity AS reserved_quantity, "
        "r.customer AS reserved_customer, r.phone AS reserved_phone, r.notes AS reservation_notes, "
        "r.created_at AS reserved_at, r.created_by AS reserved_by, si.* "
        "FROM reservations r JOIN stock_items si ON si.id = r.item_id "
        "WHERE r.status = 'active' ORDER BY r.id DESC"
    ).fetchall()
    return render_template("reservations.html", items=items, reservations=active, item_label=_item_label)


@app.route("/reservations/<int:reservation_id>/release", methods=["POST"])
@login_required
def reservation_release(reservation_id):
    db = dbmod.get_db()
    cur = db.execute(
        "UPDATE reservations SET status = 'released' WHERE id = ? AND status = 'active'",
        (reservation_id,),
    )
    if cur.rowcount:
        db.commit()
        flash("Резерв скасовано, товар знову доступний для продажу.", "success")
    else:
        flash("Активний резерв не знайдено.", "error")
    return redirect(url_for("reservations_view"))


@app.route("/reservations/<int:reservation_id>/sell", methods=["POST"])
@login_required
def reservation_sell(reservation_id):
    db = dbmod.get_db()
    reservation = db.execute(
        "SELECT si.*, r.id AS reservation_id, r.quantity AS reserved_quantity, r.customer AS reserved_customer "
        "FROM reservations r JOIN stock_items si ON si.id = r.item_id "
        "WHERE r.id = ? AND r.status = 'active'", (reservation_id,),
    ).fetchone()
    if reservation is None:
        flash("Активний резерв не знайдено.", "error")
        return redirect(url_for("reservations_view"))
    total = round(reservation["price"] * reservation["reserved_quantity"], 2)
    received = _read_received_amount(request.form.get("received", ""))
    if received is None or received < total:
        flash(f"Введіть отриману суму не меншу за {total:.2f} ₴.", "error")
        return redirect(url_for("reservations_view"))
    try:
        sale_id, _total, _received = _save_sale(
            db, [(reservation["category"], reservation, reservation["reserved_quantity"])],
            reservation["reserved_customer"], received, reservation_id=reservation_id,
        )
    except ValueError as error:
        db.rollback()
        flash(str(error), "error")
        return redirect(url_for("reservations_view"))
    db.execute(
        "UPDATE reservations SET status = 'sold', sale_id = ? WHERE id = ?",
        (sale_id, reservation["reservation_id"]),
    )
    db.commit()
    flash(f"Продаж із резерву оформлено. Решта: {received - total:.2f} ₴.", "success")
    return redirect(url_for("receipt_view", sale_id=sale_id))


@app.route("/sales/<sale_id>/cancel", methods=["POST"])
@role_required("admin")
def sale_cancel(sale_id):
    db = dbmod.get_db()
    sale_rows = db.execute(
        "SELECT * FROM history WHERE sale_id = ? AND operation = 'Продаж' ORDER BY id", (sale_id,),
    ).fetchall()
    if not sale_rows:
        flash("Продаж не знайдено.", "error")
        return redirect(url_for("dashboard"))
    sale = db.execute("SELECT reservation_id FROM sales WHERE sale_id = ?", (sale_id,)).fetchone()
    reservation_id = sale["reservation_id"] if sale else None
    db.execute(
        "INSERT OR IGNORE INTO sales (sale_id, canceled_at, canceled_by) VALUES (?, NULL, '')",
        (sale_id,),
    )
    canceled = db.execute(
        "UPDATE sales SET canceled_at = ?, canceled_by = ? WHERE sale_id = ? AND canceled_at IS NULL",
        (dbmod.now_str(), session.get("username", ""), sale_id),
    )
    if not canceled.rowcount:
        flash("Цей продаж уже скасовано.", "error")
        return redirect(url_for("receipt_view", sale_id=sale_id))

    restored_item_id = None
    for row in sale_rows:
        item_values = {field: row[field] for field in ALL_FIELDS}
        item_values.update(cost_price=row["cost_price"], price=row["price"])
        existing_item = _find_matching_item(db, row["category"], item_values)
        if existing_item:
            restored_item_id = existing_item["id"]
            db.execute(
                "UPDATE stock_items SET quantity = quantity + ?, cost_price = ?, price = ? WHERE id = ?",
                (row["quantity"], row["cost_price"], row["price"], existing_item["id"]),
            )
        else:
            fields = ", ".join(ALL_FIELDS)
            placeholders = ", ".join("?" for _ in ALL_FIELDS)
            restored = db.execute(
                f"INSERT INTO stock_items (category, {fields}, cost_price, price, quantity) "
                f"VALUES (?, {placeholders}, ?, ?, ?)",
                [row["category"]] + [item_values.get(field, "") or "" for field in ALL_FIELDS] +
                [row["cost_price"], row["price"], row["quantity"]],
            )
            restored_item_id = restored.lastrowid
        _add_history(
            db, row["category"], "Повернення", item_values, row["quantity"],
            customer=row["customer"], customer_id=row["customer_id"],
            sale_id=sale_id, username=session.get("username", ""),
        )
    if reservation_id is not None and restored_item_id is not None:
        db.execute(
            "UPDATE reservations SET item_id = ?, status = 'active' "
            "WHERE id = ? AND sale_id = ? AND status = 'sold'",
            (restored_item_id, reservation_id, sale_id),
        )
    db.commit()
    flash("Продаж скасовано, товар повернено на склад.", "success")
    return redirect(url_for("receipt_view", sale_id=sale_id))


# ------------------------------------------------------------------
#                          Чек / накладна
# ------------------------------------------------------------------

def _get_setting(db, key, default=""):
    row = db.execute("SELECT value FROM settings WHERE key = ?", (key,)).fetchone()
    if row and row["value"]:
        return row["value"]
    return default


def _render_doc_template(template_str, context):
    """Проста підстановка {{плейсхолдерів}} текстом - без виконання коду,
    тому адмін може безпечно редагувати шаблон у налаштуваннях."""
    html = template_str
    for key, value in context.items():
        replacement = "" if value is None else str(value)
        if key != "items_table":
            replacement = str(escape(replacement))
        html = html.replace("{{" + key + "}}", replacement)
    return Markup(html)


def _build_doc_context(db, rows, category, sale_id, company):
    total = sum(r["price"] * r["quantity"] for r in rows)
    total_qty = sum(r["quantity"] for r in rows)

    items_rows = []
    for i, r in enumerate(rows, start=1):
        items_rows.append(
            "<tr><td>{}</td><td>{}</td><td>{}</td><td class='nowrap'>{:.0f} ₴</td><td class='nowrap'>{:.0f} ₴</td></tr>".format(
                i, escape(_item_label(r["category"], r)), r["quantity"], r["price"], r["price"] * r["quantity"]
            )
        )
    items_table = (
        "<table class='data-table receipt-table'><thead><tr>"
        "<th>№</th><th>Позиція</th><th>К-сть</th><th>Ціна</th><th>Сума</th>"
        "</tr></thead><tbody>" + "".join(items_rows) + "</tbody></table>"
    )

    return {
        "company_name": company.get("company_name", ""),
        "company_address": company.get("company_address", ""),
        "company_phone": company.get("company_phone", ""),
        "company_edrpou": company.get("company_edrpou", ""),
        "doc_number": sale_id,
        "date": rows[0]["datetime"],
        "customer": rows[0]["customer"] or "—",
        "category_label": "Змішаний асортимент" if len({r["category"] for r in rows}) > 1 else CATEGORY_LABELS[category],
        "items_table": items_table,
        "total_qty": total_qty,
        "total_sum": f"{total:.0f} ₴",
        "amount_received": f"{total:.2f}",
        "change_amount": "0.00",
    }


@app.route("/receipt/<sale_id>")
@login_required
def receipt_view(sale_id):
    db = dbmod.get_db()
    rows = db.execute(
        "SELECT * FROM history WHERE sale_id = ? AND operation = 'Продаж' ORDER BY id", (sale_id,),
    ).fetchall()
    if not rows:
        flash("Документ не знайдено.", "error")
        return redirect(url_for("dashboard"))
    category = rows[0]["category"]
    sale = db.execute("SELECT * FROM sales WHERE sale_id = ?", (sale_id,)).fetchone()

    doc_type = request.args.get("doc", "receipt")
    if doc_type == "invoice":
        doc_type = "settlement_invoice"
    if doc_type not in ("receipt", "goods_invoice", "settlement_invoice"):
        doc_type = "receipt"

    company = {
        r["key"]: r["value"]
        for r in db.execute("SELECT key, value FROM settings WHERE key LIKE 'company_%'")
    }
    template_key, default_template = {
        "receipt": ("doc_template_receipt", DEFAULT_RECEIPT_TEMPLATE),
        "goods_invoice": ("doc_template_goods_invoice", DEFAULT_GOODS_INVOICE_TEMPLATE),
        "settlement_invoice": ("doc_template_invoice", DEFAULT_INVOICE_TEMPLATE),
    }[doc_type]
    template_str = _get_setting(db, template_key, default_template)

    context = _build_doc_context(db, rows, category, sale_id, company)
    if sale:
        context["amount_received"] = f"{sale['received']:.2f}"
        context["change_amount"] = f"{sale['change_amount']:.2f}"
    rendered = _render_doc_template(template_str, context)

    return render_template(
        "receipt.html", rendered=rendered, sale_id=sale_id, category=category, doc_type=doc_type,
        sale=sale,
    )


# ------------------------------------------------------------------
#                          База клієнтів
# ------------------------------------------------------------------

@app.route("/customers")
@login_required
def customers_view():
    db = dbmod.get_db()
    q = (request.args.get("q", "") or "").strip()
    if q:
        rows = db.execute(
            "SELECT * FROM customers WHERE name LIKE ? OR phone LIKE ? ORDER BY name",
            (f"%{q}%", f"%{q}%"),
        ).fetchall()
    else:
        rows = db.execute("SELECT * FROM customers ORDER BY name").fetchall()

    spend_rows = db.execute(
        "SELECT customer_id, SUM(price * quantity) AS total, COUNT(DISTINCT sale_id) AS sales "
        "FROM history h WHERE operation = 'Продаж' AND customer_id IS NOT NULL "
        "AND NOT EXISTS (SELECT 1 FROM sales s WHERE s.sale_id = h.sale_id AND s.canceled_at IS NOT NULL) "
        "GROUP BY customer_id"
    ).fetchall()
    spend = {r["customer_id"]: {"total": r["total"], "sales": r["sales"]} for r in spend_rows}

    return render_template("customers.html", customers=rows, q=q, spend=spend)


@app.route("/customers/add", methods=["POST"])
@login_required
def customers_add():
    db = dbmod.get_db()
    name = (request.form.get("name", "") or "").strip()
    phone = (request.form.get("phone", "") or "").strip()
    notes = (request.form.get("notes", "") or "").strip()
    if not name:
        flash("Вкажіть ім'я клієнта.", "error")
        return redirect(url_for("customers_view"))
    existing = db.execute("SELECT id FROM customers WHERE LOWER(name) = LOWER(?)", (name,)).fetchone()
    if existing:
        flash("Клієнт з таким іменем вже є в базі.", "error")
        return redirect(url_for("customer_detail", customer_id=existing["id"]))
    cur = db.execute(
        "INSERT INTO customers (name, phone, notes, created_at) VALUES (?, ?, ?, ?)",
        (name, phone, notes, dbmod.now_str()),
    )
    db.commit()
    flash(f"Клієнта «{name}» додано.", "success")
    return redirect(url_for("customer_detail", customer_id=cur.lastrowid))


@app.route("/customers/<int:customer_id>")
@login_required
def customer_detail(customer_id):
    db = dbmod.get_db()
    customer = db.execute("SELECT * FROM customers WHERE id = ?", (customer_id,)).fetchone()
    if not customer:
        flash("Клієнта не знайдено.", "error")
        return redirect(url_for("customers_view"))
    purchases = db.execute(
        "SELECT * FROM history h WHERE customer_id = ? AND operation = 'Продаж' "
        "AND NOT EXISTS (SELECT 1 FROM sales s WHERE s.sale_id = h.sale_id AND s.canceled_at IS NOT NULL) "
        "ORDER BY id DESC",
        (customer_id,),
    ).fetchall()
    total_spent = sum(r["price"] * r["quantity"] for r in purchases)
    return render_template(
        "customer_detail.html", customer=customer, purchases=purchases,
        total_spent=total_spent, item_label=_item_label,
    )


@app.route("/customers/edit/<int:customer_id>", methods=["POST"])
@login_required
def customers_edit(customer_id):
    db = dbmod.get_db()
    name = (request.form.get("name", "") or "").strip()
    phone = (request.form.get("phone", "") or "").strip()
    notes = (request.form.get("notes", "") or "").strip()
    if not name:
        flash("Ім'я не може бути порожнім.", "error")
        return redirect(url_for("customer_detail", customer_id=customer_id))
    db.execute(
        "UPDATE customers SET name = ?, phone = ?, notes = ? WHERE id = ?",
        (name, phone, notes, customer_id),
    )
    db.commit()
    flash("Дані клієнта оновлено.", "success")
    return redirect(url_for("customer_detail", customer_id=customer_id))


@app.route("/customers/delete/<int:customer_id>", methods=["POST"])
@role_required("admin")
def customers_delete(customer_id):
    db = dbmod.get_db()
    db.execute("DELETE FROM customers WHERE id = ?", (customer_id,))
    db.commit()
    flash("Клієнта видалено (історія його покупок залишається в записах продажів).", "success")
    return redirect(url_for("customers_view"))


# ------------------------------------------------------------------
#                     Масовий імпорт (CSV / Excel)
# ------------------------------------------------------------------

def _parse_import_file(file_storage):
    """Читає завантажений CSV або XLSX у список словників {заголовок: значення}."""
    filename = file_storage.filename or ""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext == "csv":
        raw = file_storage.read()
        text = raw.decode("utf-8-sig", errors="replace")
        sample = text[:2048]
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
        reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
        rows = [{(k or "").strip(): (v or "").strip() for k, v in row.items()} for row in reader]
        return rows, None

    if ext in ("xlsx", "xlsm"):
        if openpyxl is None:
            return None, "Підтримка Excel не встановлена на сервері. Скористайтесь форматом CSV."
        wb = openpyxl.load_workbook(file_storage, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            return [], None
        header = [str(h).strip() if h is not None else "" for h in header]
        rows = []
        for raw_row in rows_iter:
            row = {}
            for i, val in enumerate(raw_row):
                if i < len(header) and header[i]:
                    row[header[i]] = "" if val is None else str(val).strip()
            if any(v for v in row.values()):
                rows.append(row)
        return rows, None

    return None, "Підтримувані формати: CSV або Excel (.xlsx)."


@app.route("/stock/<category>/import", methods=["GET", "POST"])
@role_required("admin")
def item_import(category):
    _category_or_404(category)
    cat_type = CATEGORY_TYPE[category]
    columns = IMPORT_COLUMNS[cat_type]

    if request.method == "POST":
        file = request.files.get("file")
        if not file or file.filename == "":
            flash("Виберіть файл CSV або Excel.", "error")
            return redirect(url_for("item_import", category=category))

        rows, error = _parse_import_file(file)
        if error:
            flash(error, "error")
            return redirect(url_for("item_import", category=category))
        if not rows:
            flash("У файлі не знайдено жодного рядка з даними.", "error")
            return redirect(url_for("item_import", category=category))

        header_map = {h.lower(): key for h, key in columns}
        valid_color_codes = {code for code, _label, _hex in COLOR_TAGS}

        db = dbmod.get_db()
        added, merged, skipped = 0, 0, 0
        errors = []

        for i, raw_row in enumerate(rows, start=2):  # рядок 1 у файлі - заголовок
            norm = {}
            for h, v in raw_row.items():
                key = header_map.get((h or "").strip().lower())
                if key:
                    norm[key] = (v or "").strip()

            brand = norm.get("brand", "").strip()
            if not brand:
                skipped += 1
                if len(errors) < 20:
                    errors.append(f"Рядок {i}: не вказано бренд - пропущено.")
                continue

            try:
                quantity = int(float((norm.get("quantity", "0") or "0").replace(",", ".")))
            except ValueError:
                quantity = 0
            if quantity <= 0:
                skipped += 1
                if len(errors) < 20:
                    errors.append(f"Рядок {i} ({brand}): некоректна або нульова кількість - пропущено.")
                continue

            try:
                price = float((norm.get("price", "0") or "0").replace(",", "."))
            except ValueError:
                price = 0.0
            try:
                cost_price = float((norm.get("cost_price", "0") or "0").replace(",", "."))
            except ValueError:
                cost_price = 0.0

            color_tag = norm.get("color_tag", "").strip()
            if color_tag not in valid_color_codes and not color_tag.startswith("#"):
                color_tag = ""

            values = {f: norm.get(f, "") for f in ALL_FIELDS}
            if cat_type == "battery":
                for numf in ("capacity_ah", "cold_crank_amps"):
                    try:
                        values[numf] = float((values.get(numf) or "0").replace(",", "."))
                    except ValueError:
                        values[numf] = 0
            if cat_type == "tire" and values.get("dot_code") and not is_valid_dot_code(values["dot_code"]):
                values["dot_code"] = ""

            existing = _find_matching_item(db, category, values)
            if existing:
                db.execute(
                    "UPDATE stock_items SET quantity = quantity + ?, price = ?, cost_price = ? WHERE id = ?",
                    (quantity, price, cost_price, existing["id"]),
                )
                merged += 1
            else:
                fields = ", ".join(ALL_FIELDS)
                placeholders = ", ".join("?" for _ in ALL_FIELDS)
                vals = [values.get(f, "") for f in ALL_FIELDS]
                db.execute(
                    f"INSERT INTO stock_items (category, {fields}, color_tag, cost_price, price, quantity) "
                    f"VALUES (?, {placeholders}, ?, ?, ?, ?)",
                    [category] + vals + [color_tag, cost_price, price, quantity],
                )
                added += 1

            values["price"] = price
            values["cost_price"] = cost_price
            values["color_tag"] = color_tag
            _add_history(db, category, "Імпорт", values, quantity)

        db.commit()
        flash(f"Імпорт завершено: {added} нових, {merged} об'єднано з наявними, {skipped} пропущено.",
              "success" if (added or merged) else "error")
        for e in errors:
            flash(e, "error")
        return redirect(url_for("stock_list", category=category))

    return render_template("import.html", category=category, columns=columns)


@app.route("/stock/<category>/import/template.csv")
@role_required("admin")
def import_template_csv(category):
    _category_or_404(category)
    cat_type = CATEGORY_TYPE[category]
    header = [h for h, _key in IMPORT_COLUMNS[cat_type]]
    return _csv_response([], header, f"{category}_shablon.csv")


# ------------------------------------------------------------------
#                          Історія
# ------------------------------------------------------------------

@app.route("/stock/<category>/history")
@login_required
def history_view(category):
    _category_or_404(category)
    db = dbmod.get_db()
    rows = db.execute(
        "SELECT * FROM history WHERE category = ? ORDER BY id DESC LIMIT 300", (category,)
    ).fetchall()
    return render_template("history.html", category=category, rows=rows)


# ------------------------------------------------------------------
#                          Налаштування (зміна пароля)
# ------------------------------------------------------------------

@app.route("/settings", methods=["GET", "POST"])
@login_required
def settings_view():
    db = dbmod.get_db()
    if request.method == "POST":
        form_type = request.form.get("form_type", "password")

        if form_type == "company" and session.get("role") == "admin":
            for key in ("company_name", "company_phone", "company_address", "company_edrpou"):
                value = (request.form.get(key, "") or "").strip()
                db.execute(
                    "INSERT INTO settings (key, value) VALUES (?, ?) "
                    "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
                    (key, value),
                )
            db.commit()
            flash("Реквізити для чеків/накладних збережено.", "success")

        elif form_type == "templates" and session.get("role") == "admin":
            for key in ("doc_template_receipt", "doc_template_goods_invoice", "doc_template_invoice"):
                if request.form.get(f"reset_{key}"):
                    db.execute("DELETE FROM settings WHERE key = ?", (key,))
                else:
                    value = request.form.get(key, "")
                    db.execute(
                        "INSERT INTO settings (key, value) VALUES (?, ?) "
                        "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
                        (key, value),
                    )
            db.commit()
            flash("Шаблони друку збережено.", "success")

        else:
            old_pw = request.form.get("old_password", "")
            new_pw = request.form.get("new_password", "")
            new_pw2 = request.form.get("new_password2", "")
            user = db.execute("SELECT * FROM users WHERE id = ?", (session.get("user_id"),)).fetchone()

            if not user or not check_password_hash(user["password_hash"], old_pw):
                flash("Поточний пароль невірний.", "error")
            elif len(new_pw) < 4:
                flash("Новий пароль має бути не коротшим за 4 символи.", "error")
            elif new_pw != new_pw2:
                flash("Новий пароль і підтвердження не збігаються.", "error")
            else:
                db.execute(
                    "UPDATE users SET password_hash = ? WHERE id = ?",
                    (generate_password_hash(new_pw), user["id"]),
                )
                db.commit()
                flash("Пароль змінено.", "success")

    company = {}
    templates = {}
    if session.get("role") == "admin":
        rows = db.execute("SELECT key, value FROM settings WHERE key LIKE 'company_%'").fetchall()
        company = {r["key"]: r["value"] for r in rows}
        templates["doc_template_receipt"] = _get_setting(db, "doc_template_receipt", DEFAULT_RECEIPT_TEMPLATE)
        templates["doc_template_goods_invoice"] = _get_setting(
            db, "doc_template_goods_invoice", DEFAULT_GOODS_INVOICE_TEMPLATE,
        )
        templates["doc_template_invoice"] = _get_setting(db, "doc_template_invoice", DEFAULT_INVOICE_TEMPLATE)
    return render_template("settings.html", company=company, templates=templates, placeholders=DOC_PLACEHOLDERS_HELP)


# ------------------------------------------------------------------
#                     Користувачі (тільки адмін)
# ------------------------------------------------------------------

@app.route("/users")
@role_required("admin")
def users_view():
    db = dbmod.get_db()
    users = db.execute("SELECT * FROM users ORDER BY id").fetchall()
    admin_count = sum(1 for u in users if u["role"] == "admin")
    return render_template("users.html", users=users, admin_count=admin_count)


@app.route("/users/add", methods=["POST"])
@role_required("admin")
def users_add():
    db = dbmod.get_db()
    username = (request.form.get("username", "") or "").strip()
    password = request.form.get("password", "")
    role = request.form.get("role", "seller")
    if role not in ("admin", "seller"):
        role = "seller"

    if not username or len(username) < 2:
        flash("Логін має бути не коротшим за 2 символи.", "error")
    elif len(password) < 4:
        flash("Пароль має бути не коротшим за 4 символи.", "error")
    elif db.execute("SELECT 1 FROM users WHERE username = ?", (username,)).fetchone():
        flash("Такий логін вже існує.", "error")
    else:
        db.execute(
            "INSERT INTO users (username, password_hash, role, created_at) VALUES (?, ?, ?, ?)",
            (username, generate_password_hash(password), role, dbmod.now_str()),
        )
        db.commit()
        flash(f"Користувача «{username}» створено.", "success")
    return redirect(url_for("users_view"))


@app.route("/users/delete/<int:user_id>", methods=["POST"])
@role_required("admin")
def users_delete(user_id):
    db = dbmod.get_db()
    if user_id == session.get("user_id"):
        flash("Не можна видалити самого себе.", "error")
        return redirect(url_for("users_view"))

    target = db.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
    if not target:
        flash("Користувача не знайдено.", "error")
        return redirect(url_for("users_view"))

    if target["role"] == "admin":
        admin_count = db.execute("SELECT COUNT(*) AS c FROM users WHERE role = 'admin'").fetchone()["c"]
        if admin_count <= 1:
            flash("Не можна видалити останнього адміністратора.", "error")
            return redirect(url_for("users_view"))

    db.execute("DELETE FROM users WHERE id = ?", (user_id,))
    db.commit()
    flash(f"Користувача «{target['username']}» видалено.", "success")
    return redirect(url_for("users_view"))


@app.route("/users/reset/<int:user_id>", methods=["POST"])
@role_required("admin")
def users_reset_password(user_id):
    db = dbmod.get_db()
    new_pw = request.form.get("new_password", "")
    if len(new_pw) < 4:
        flash("Пароль має бути не коротшим за 4 символи.", "error")
        return redirect(url_for("users_view"))
    user = db.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
    if not user:
        flash("Користувача не знайдено.", "error")
        return redirect(url_for("users_view"))
    db.execute("UPDATE users SET password_hash = ? WHERE id = ?", (generate_password_hash(new_pw), user_id))
    db.commit()
    flash(f"Пароль для «{user['username']}» скинуто.", "success")
    return redirect(url_for("users_view"))


# ------------------------------------------------------------------
#                          Експорт у CSV
# ------------------------------------------------------------------

def _csv_response(rows, header, filename):
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    writer.writerow(header)
    writer.writerows(rows)
    output = "\ufeff" + buf.getvalue()  # BOM, щоб Excel коректно показував українські літери
    return Response(
        output, mimetype="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.route("/stock/<category>/export.csv")
@login_required
def export_stock_csv(category):
    _category_or_404(category)
    cat_type = CATEGORY_TYPE[category]
    db = dbmod.get_db()
    items = db.execute("SELECT * FROM stock_items WHERE category = ? ORDER BY brand, size, title", (category,)).fetchall()
    is_admin = session.get("role") == "admin"

    def cost_cols(i):
        if not is_admin:
            return []
        margin = i["price"] - i["cost_price"]
        return [i["cost_price"], margin]

    cost_header = ["Собівартість", "Маржа"] if is_admin else []

    if cat_type == "tire":
        header = ["Бренд", "Модель", "Розмір", "Сезон", "LI", "SI", "DOT", "Мітка"] + cost_header + ["Ціна", "Кількість"]
        rows = [[i["brand"], i["tread_model"], i["size"], i["season"], i["load_index"], i["speed_index"],
                 i["dot_code"], i["color_tag"]] + cost_cols(i) + [i["price"], i["quantity"]] for i in items]
    elif cat_type == "battery":
        header = ["Бренд", "Тип", "Ємність (Ah)", "Пусковий струм (A)", "Полярність", "Корпус", "Мітка"] + cost_header + ["Ціна", "Кількість"]
        rows = [[i["brand"], i["transport_type"], i["capacity_ah"], i["cold_crank_amps"], i["polarity"],
                 i["case_type"], i["color_tag"]] + cost_cols(i) + [i["price"], i["quantity"]] for i in items]
    else:
        header = ["Бренд", "Назва", "В'язкість", "Основа", "Обʼєм", "Мітка"] + cost_header + ["Ціна", "Кількість"]
        rows = [[i["brand"], i["title"], i["viscosity"], i["base_type"], i["volume"],
                 i["color_tag"]] + cost_cols(i) + [i["price"], i["quantity"]] for i in items]

    return _csv_response(rows, header, f"{category}_sklad.csv")


@app.route("/stock/<category>/history/export.csv")
@login_required
def export_history_csv(category):
    _category_or_404(category)
    cat_type = CATEGORY_TYPE[category]
    db = dbmod.get_db()
    rows_db = db.execute("SELECT * FROM history WHERE category = ? ORDER BY id DESC", (category,)).fetchall()
    is_admin = session.get("role") == "admin"

    def cost_cols(r):
        if not is_admin:
            return []
        margin = (r["price"] - r["cost_price"]) * r["quantity"] if r["operation"] == "Продаж" else 0
        return [r["cost_price"], margin]

    cost_header = ["Собівартість", "Прибуток"] if is_admin else []

    if cat_type == "tire":
        header = ["Коли", "Дія", "Хто", "Бренд", "Розмір", "Сезон", "DOT"] + cost_header + ["Кількість", "Ціна", "Клієнт"]
        rows = [[r["datetime"], r["operation"], r["username"], r["brand"], r["size"], r["season"], r["dot_code"]] +
                cost_cols(r) + [r["quantity"], r["price"], r["customer"]] for r in rows_db]
    elif cat_type == "battery":
        header = ["Коли", "Дія", "Хто", "Бренд", "Тип", "Ємність (Ah)", "Пусковий струм (A)"] + cost_header + ["Кількість", "Ціна", "Клієнт"]
        rows = [[r["datetime"], r["operation"], r["username"], r["brand"], r["transport_type"], r["capacity_ah"],
                 r["cold_crank_amps"]] + cost_cols(r) + [r["quantity"], r["price"], r["customer"]] for r in rows_db]
    else:
        header = ["Коли", "Дія", "Хто", "Бренд", "Назва", "В'язкість", "Обʼєм"] + cost_header + ["Кількість", "Ціна", "Клієнт"]
        rows = [[r["datetime"], r["operation"], r["username"], r["brand"], r["title"], r["viscosity"], r["volume"]] +
                cost_cols(r) + [r["quantity"], r["price"], r["customer"]] for r in rows_db]

    return _csv_response(rows, header, f"{category}_istoriya.csv")


# ------------------------------------------------------------------
#                          Іконки брендів
# ------------------------------------------------------------------

@app.route("/brand-icons")
@role_required("admin")
def brand_icons_view():
    db = dbmod.get_db()
    custom_brands = {r["brand"] for r in db.execute("SELECT DISTINCT brand FROM stock_items WHERE brand != ''")}
    all_brands = sorted(set(ALL_KNOWN_BRANDS) | custom_brands)
    icons = _get_brand_icons(db)
    return render_template("brand_icons.html", all_brands=all_brands, icons=icons)


@app.route("/brand-icons/upload", methods=["POST"])
@role_required("admin")
def brand_icon_upload():
    brand = (request.form.get("brand", "") or "").strip()
    file = request.files.get("file")
    if not brand:
        flash("Не вказано бренд.", "error")
        return redirect(url_for("brand_icons_view"))
    if not file or file.filename == "":
        flash("Виберіть файл PNG або ICO.", "error")
        return redirect(url_for("brand_icons_view"))

    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in ALLOWED_ICON_EXT:
        flash("Дозволені формати іконок: PNG або ICO.", "error")
        return redirect(url_for("brand_icons_view"))

    db = dbmod.get_db()
    old = db.execute("SELECT icon_file FROM brand_icons WHERE brand = ?", (brand,)).fetchone()
    if old and old["icon_file"]:
        old_path = os.path.join(BRAND_ICON_DIR, old["icon_file"])
        if os.path.exists(old_path):
            os.remove(old_path)

    filename = f"{_slugify(brand)}-{secrets.token_hex(4)}.{ext}"
    file.save(os.path.join(BRAND_ICON_DIR, filename))

    db.execute(
        "INSERT INTO brand_icons (brand, icon, icon_file) VALUES (?, '', ?) "
        "ON CONFLICT(brand) DO UPDATE SET icon_file = excluded.icon_file, icon = ''",
        (brand, filename),
    )
    db.commit()
    flash(f"Іконку для «{brand}» завантажено.", "success")
    return redirect(url_for("brand_icons_view"))


@app.route("/brand-icons/set-emoji", methods=["POST"])
@role_required("admin")
def brand_icon_set_emoji():
    brand = (request.form.get("brand", "") or "").strip()
    icon = (request.form.get("icon", "") or "").strip()
    if not brand:
        return redirect(url_for("brand_icons_view"))
    db = dbmod.get_db()
    if icon:
        old = db.execute("SELECT icon_file FROM brand_icons WHERE brand = ?", (brand,)).fetchone()
        if old and old["icon_file"]:
            old_path = os.path.join(BRAND_ICON_DIR, old["icon_file"])
            if os.path.exists(old_path):
                os.remove(old_path)
        db.execute(
            "INSERT INTO brand_icons (brand, icon, icon_file) VALUES (?, ?, '') "
            "ON CONFLICT(brand) DO UPDATE SET icon = excluded.icon, icon_file = ''",
            (brand, icon),
        )
        db.commit()
    else:
        db.execute("UPDATE brand_icons SET icon = '' WHERE brand = ?", (brand,))
        db.commit()
    flash(f"Іконку для «{brand}» збережено.", "success")
    return redirect(url_for("brand_icons_view"))


@app.route("/brand-icons/remove", methods=["POST"])
@role_required("admin")
def brand_icon_remove():
    brand = (request.form.get("brand", "") or "").strip()
    db = dbmod.get_db()
    row = db.execute("SELECT icon_file FROM brand_icons WHERE brand = ?", (brand,)).fetchone()
    if row and row["icon_file"]:
        path = os.path.join(BRAND_ICON_DIR, row["icon_file"])
        if os.path.exists(path):
            os.remove(path)
    db.execute("DELETE FROM brand_icons WHERE brand = ?", (brand,))
    db.commit()
    flash(f"Іконку для «{brand}» видалено.", "success")
    return redirect(url_for("brand_icons_view"))


# ------------------------------------------------------------------
#                     Калькулятор перевзування
# ------------------------------------------------------------------

@app.route("/calculator")
@login_required
def calculator_view():
    result = None
    form = {k: request.args.get(k, "") for k in ("w1", "p1", "d1", "w2", "p2", "d2")}
    if all(form.values()):
        try:
            result = compare_tire_sizes(
                float(form["w1"]), float(form["p1"]), float(form["d1"]),
                float(form["w2"]), float(form["p2"]), float(form["d2"]),
            )
        except ValueError:
            flash("Некоректні значення розмірів.", "error")
    return render_template(
        "calculator.html", form=form, result=result,
        car_widths=CAR_WIDTHS, car_profiles=CAR_PROFILES,
        car_diameter_inches=[int(d[1:]) for d in CAR_DIAMETERS],
    )


if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5000)
